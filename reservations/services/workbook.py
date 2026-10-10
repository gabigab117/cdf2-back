"""The reservations of an event as an Excel workbook, in the v1's format."""

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from events.models import Event

# The type of an Excel workbook, as the export declares it.
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def reservations_workbook(event: Event) -> bytes:
    """One sheet, « Réservations »: a row per reservation in the order they
    were recorded, a column per type of place, and a row of totals.

    The headers and the totals are bold, the first row stays in view, and the
    dates read in Paris, without a time zone: Excel knows none.
    """
    ticket_types = list(event.ticket_types.all())
    headers = [
        "#",
        "Nom",
        "Remarque",
        "Saisie le",
        *(ticket_type.name for ticket_type in ticket_types),
        "Total",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Réservations"
    sheet.append(headers)
    totals = [0] * len(ticket_types)
    reservations = event.reservations.prefetch_related("lines").order_by("created_at", "pk")
    for number, reservation in enumerate(reservations, start=1):
        quantities = {line.ticket_type_id: line.quantity for line in reservation.lines.all()}
        row = [quantities.get(ticket_type.pk, 0) for ticket_type in ticket_types]
        totals = [total + quantity for total, quantity in zip(totals, row, strict=True)]
        recorded = timezone.localtime(reservation.created_at).replace(tzinfo=None)
        sheet.append([number, reservation.name, reservation.note, recorded, *row, sum(row)])
        sheet.cell(row=sheet.max_row, column=4).number_format = "DD/MM/YYYY HH:MM"
    sheet.append(["Total", None, None, None, *totals, sum(totals)])
    bold = Font(bold=True)
    for cell in (*sheet[1], *sheet[sheet.max_row]):
        cell.font = bold
    sheet.freeze_panes = "A2"
    for column, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = max(12, len(header) + 2)
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()

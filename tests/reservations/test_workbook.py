import datetime as dt
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from reservations.models import Reservation
from tests.events.factories import EventFactory
from tests.reservations.factories import TicketTypeFactory, reserve

pytestmark = pytest.mark.django_db

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_url(event):
    return f"/api/board/events/{event.id}/reservations.xlsx"


def sheet_of(response):
    return load_workbook(BytesIO(response.content)).active


def recorded_at(reservation, moment):
    Reservation.objects.filter(pk=reservation.pk).update(created_at=moment)


@pytest.fixture
def meal():
    """« Repas des aînés », with its two menus and two reservations recorded on
    3 October 2026 at 9:30 and 14:05 in Paris.
    """
    event = EventFactory(title="Repas des aînés", slug="repas-des-aines-2026")
    adult = TicketTypeFactory(event=event, name="Menu adulte", sort_order=0)
    child = TicketTypeFactory(event=event, name="Menu enfant", sort_order=1)
    martin = reserve(
        event, name="Famille Martin", note="Table 4", adulte=(adult, 2), enfant=(child, 4)
    )
    petit = reserve(event, name="Famille Petit", adulte=(adult, 1))
    recorded_at(martin, dt.datetime(2026, 10, 3, 7, 30, tzinfo=dt.UTC))
    recorded_at(petit, dt.datetime(2026, 10, 3, 12, 5, tzinfo=dt.UTC))
    return event


def test_the_export_is_an_excel_file_named_after_the_event(board_client, meal):
    """
    Given the reservations of « Repas des aînés »
    When the board exports them
    Then it receives an Excel workbook, as a file named after the event's address
    """
    response = board_client.get(export_url(meal))

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX
    assert (
        response["Content-Disposition"]
        == 'attachment; filename="reservations-repas-des-aines-2026.xlsx"'
    )


def test_the_export_holds_a_row_per_reservation_and_their_totals(board_client, meal):
    """
    Given two reservations of the meal, recorded at 9:30 then at 14:05 in Paris
    When the board exports them
    Then the sheet « Réservations » holds the headers, a row per reservation in
    the order they were recorded, and a row of totals
    And the dates read in Paris, without a time zone, as JJ/MM/AAAA HH:MM
    """
    sheet = sheet_of(board_client.get(export_url(meal)))

    assert sheet.title == "Réservations"
    assert [list(row) for row in sheet.iter_rows(values_only=True)] == [
        ["#", "Nom", "Remarque", "Saisie le", "Menu adulte", "Menu enfant", "Total"],
        [1, "Famille Martin", "Table 4", dt.datetime(2026, 10, 3, 9, 30), 2, 4, 6],
        [2, "Famille Petit", None, dt.datetime(2026, 10, 3, 14, 5), 1, 0, 1],
        ["Total", None, None, None, 3, 4, 7],
    ]
    assert sheet["D2"].number_format == "DD/MM/YYYY HH:MM"


def test_the_headers_and_totals_are_bold_and_the_headers_stay_in_view(board_client, meal):
    """
    Given the reservations of the meal
    When the board exports them
    Then the headers and the totals are bold, the first row stays in view, and
    every column is wide enough for its header
    """
    sheet = sheet_of(board_client.get(export_url(meal)))

    assert all(cell.font.bold for cell in sheet[1])
    assert all(cell.font.bold for cell in sheet[sheet.max_row])
    assert not sheet["A2"].font.bold
    assert sheet.freeze_panes == "A2"
    assert sheet.column_dimensions["A"].width == 12
    assert sheet.column_dimensions["E"].width == len("Menu adulte") + 2


def test_an_event_without_reservations_exports_its_headers_and_zero_totals(board_client):
    """
    Given an event with a type of place and no reservation
    When the board exports its reservations
    Then the sheet holds the headers and a row of zero totals
    """
    event = EventFactory()
    TicketTypeFactory(event=event, name="Entrée")

    rows = list(sheet_of(board_client.get(export_url(event))).iter_rows(values_only=True))

    assert rows == [
        ("#", "Nom", "Remarque", "Saisie le", "Entrée", "Total"),
        ("Total", None, None, None, 0, 0),
    ]


def test_an_event_without_types_exports_the_columns_of_every_reservation(board_client):
    """
    Given an event without types of place
    When the board exports its reservations
    Then the headers hold the columns every reservation has
    """
    event = EventFactory()

    rows = list(sheet_of(board_client.get(export_url(event))).iter_rows(values_only=True))

    assert rows[0] == ("#", "Nom", "Remarque", "Saisie le", "Total")


def test_the_dates_of_the_export_follow_the_summer_time_of_paris(board_client):
    """
    Given a reservation recorded on 15 July 2026 at 10:00 UTC
    When the board exports it
    Then its date reads 12:00, in Paris in summer
    """
    event = EventFactory()
    menu = TicketTypeFactory(event=event)
    reservation = reserve(event, menu=(menu, 1))
    recorded_at(reservation, dt.datetime(2026, 7, 15, 10, tzinfo=dt.UTC))

    sheet = sheet_of(board_client.get(export_url(event)))

    assert sheet["D2"].value == dt.datetime(2026, 7, 15, 12)
    assert timezone.get_current_timezone_name() == "Europe/Paris"
